"""Readers for the phase-1 file types: CSV, Excel, Word.

Each reader returns a list of *tables*::

    [{"name": "Sheet1", "kind": "table", "frame": <DataFrame>}, ...]
    [{"name": "document", "kind": "document", "text": "...", "headings": [...]}, ...]
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd

CSV_EXTS = {".csv"}
EXCEL_EXTS = {".xlsx", ".xls"}
WORD_EXTS = {".docx"}


def detect_file_type(path: Path) -> str | None:
    ext = path.suffix.lower()
    if ext in CSV_EXTS:
        return "csv"
    if ext in EXCEL_EXTS:
        return "excel"
    if ext in WORD_EXTS:
        return "word"
    return None


def read_csv(path: Path) -> list[dict[str, Any]]:
    # Try utf-8 first, fall back to latin-1 — real-world team files vary.
    for enc in ("utf-8", "latin-1"):
        try:
            df = pd.read_csv(path, encoding=enc)
            return [{"name": "data", "kind": "table", "frame": df}]
        except (UnicodeDecodeError, pd.errors.ParserError):
            continue
    # Last resort: let pandas figure it out, replacing bad bytes.
    df = pd.read_csv(path, encoding="utf-8", encoding_errors="replace")
    return [{"name": "data", "kind": "table", "frame": df}]


def read_excel(path: Path) -> list[dict[str, Any]]:
    sheets: dict[str, pd.DataFrame] = pd.read_excel(path, sheet_name=None)
    tables = []
    for name, df in sheets.items():
        # Skip sheets that are entirely empty.
        if df.dropna(how="all").empty:
            continue
        tables.append({"name": str(name), "kind": "table", "frame": df})
    if not tables:
        raise ValueError(f"Workbook {path.name} contains no non-empty sheets")
    return tables


def read_word(path: Path) -> list[dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    headings = [
        p.text.strip()
        for p in doc.paragraphs
        if p.style and p.style.name.startswith("Heading") and p.text.strip()
    ]
    text = "\n\n".join(paragraphs)

    tables: list[dict[str, Any]] = [
        {
            "name": "document",
            "kind": "document",
            "text": text,
            "headings": headings,
            "n_paragraphs": len(paragraphs),
            "n_tables": len(doc.tables),
        }
    ]
    # Word tables become proper tabular entries too, so LLMs can query them.
    for i, tbl in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
        if not rows:
            continue
        header, *body = rows
        df = pd.DataFrame(body, columns=[h or f"col_{j}" for j, h in enumerate(header)])
        tables.append({"name": f"word_table_{i}", "kind": "table", "frame": df})
    return tables


def read_file(path: Path) -> list[dict[str, Any]]:
    ftype = detect_file_type(path)
    if ftype == "csv":
        return read_csv(path)
    if ftype == "excel":
        return read_excel(path)
    if ftype == "word":
        return read_word(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")
