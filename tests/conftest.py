import shutil
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from hub.config import HubConfig  # noqa: E402


@pytest.fixture()
def cfg(tmp_path) -> HubConfig:
    """A HubConfig rooted entirely inside a tmp dir — never touches data/."""
    raw = {
        "hub": {"name": "Test Hub"},
        "paths": {
            "inbox": str(tmp_path / "inbox"),
            "processed": str(tmp_path / "processed"),
            "failed": str(tmp_path / "failed"),
            "catalog_db": str(tmp_path / "catalog" / "catalog.db"),
            "catalog_json": str(tmp_path / "catalog" / "assets.json"),
            "sources_dir": str(tmp_path / "sources"),
        },
        "ingest": {
            "supported_extensions": [".csv", ".xlsx", ".xls", ".docx"],
            "settle_seconds": 0,
            "column_sample_values": 3,
            "doc_text_chars": 5000,
            "after_ingest": "move",
        },
    }
    config = HubConfig(raw, tmp_path)
    config.ensure_dirs()
    return config


@pytest.fixture()
def sample_csv(tmp_path) -> Path:
    p = tmp_path / "yield_trials_2025.csv"
    p.write_text(
        "region,crop,yield_kg_ha\n"
        "Oromia,sorghum,2450\n"
        "Amhara,teff,1830\n"
        "SNNP,maize,3100\n",
        encoding="utf-8",
    )
    return p


@pytest.fixture()
def sample_xlsx(tmp_path) -> Path:
    import pandas as pd

    p = tmp_path / "soil_samples.xlsx"
    with pd.ExcelWriter(p) as xw:
        pd.DataFrame({"plot": ["A1", "A2"], "ph": [5.6, 6.1]}).to_excel(
            xw, sheet_name="Soil pH", index=False
        )
        pd.DataFrame({"plot": ["A1", "A2"], "oc_pct": [1.2, 1.5]}).to_excel(
            xw, sheet_name="Organic Carbon", index=False
        )
    return p


@pytest.fixture()
def sample_docx(tmp_path) -> Path:
    from docx import Document

    p = tmp_path / "advisory_protocol.docx"
    doc = Document()
    doc.add_heading("Fertilizer Advisory Protocol", level=1)
    doc.add_paragraph("This protocol describes woreda-level advisory generation.")
    doc.add_paragraph("Inputs: soil rasters, AEZ boundaries, yield models.")
    doc.save(p)
    return p
