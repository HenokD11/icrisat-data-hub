"""Build the ICRISAT Data Hub report deck from the Hyper-Localized Fertilizer
Advisory template — same master, decorations, colours, and slide patterns.

Run:  .venv\\Scripts\\python sample_report\\build_report.py
Out:  sample_report\\ICRISAT_Data_Hub_Report.pptx
"""

import copy
import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

BASE = Path(__file__).resolve().parent
SRC = BASE / "Hyper_Localized_Fertilizer_Advisory.pptx"
OUT = BASE / "ICRISAT_Data_Hub_Report.pptx"
SHOTS = BASE / "screenshots"

prs = Presentation(str(SRC))
N_ORIG = len(prs.slides._sldIdLst)

# ---------------------------------------------------------------- duplication
def _collect_pictures(shapes, ox=0, oy=0, sx=1.0, sy=1.0):
    """Recursively capture picture blobs + absolute slide positions from a slide."""
    pics = []
    for shape in shapes:
        if shape.shape_type == 6:  # group: children use child coordinate space
            try:
                xfrm = shape._element.grpSpPr.xfrm
                chox, choy = xfrm.chOff.x, xfrm.chOff.y
                chcx, chcy = xfrm.chExt.cx, xfrm.chExt.cy
                gsx = shape.width / chcx if chcx else 1.0
                gsy = shape.height / chcy if chcy else 1.0
                gx = ox + shape.left * sx
                gy = oy + shape.top * sy
                pics += _collect_pictures(shape.shapes, gx - chox * gsx * sx, gy - choy * gsy * sy, sx * gsx, sy * gsy)
            except Exception:
                pass
        elif shape.shape_type == 13:
            try:
                blob = shape.image.blob
            except Exception:
                continue
            pics.append({
                "name": shape.name,
                "blob": blob,
                "left": int(ox + shape.left * sx),
                "top": int(oy + shape.top * sy),
                "width": int(shape.width * sx),
                "height": int(shape.height * sy),
            })
    return pics


def dup_slide(idx):
    """Deep-copy a template slide; returns (dest_slide, captured_pictures)."""
    source = prs.slides[idx]
    pics = _collect_pictures(source.shapes)
    dest = prs.slides.add_slide(source.slide_layout)
    for shape in source.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shape.element))
    return dest, pics


def _fit(width, height, aspect):
    """Fit an image of given aspect into a slot, centred; returns (dl,dt,w,h)."""
    slot_aspect = width / height
    if aspect >= slot_aspect:
        w, h = width, int(width / aspect)
    else:
        w, h = int(height * aspect), height
    return (width - w) // 2, (height - h) // 2, w, h


def fix_pictures(slide, pics, replacements=None):
    """Drop every picture on the duplicated slide (broken refs) and re-add
    from the source-captured blobs, or from replacement file paths."""
    replacements = replacements or {}
    for shape in list(slide.shapes):
        _drop_pictures(shape)
    for pic in pics:
        path = replacements.get(pic["name"])
        data = Path(path).read_bytes() if path else pic["blob"]
        aspect = Image.open(io.BytesIO(data)).width / Image.open(io.BytesIO(data)).height
        dl, dt, w, h = _fit(pic["width"], pic["height"], aspect)
        slide.shapes.add_picture(io.BytesIO(data), pic["left"] + dl, pic["top"] + dt, w, h)


def _drop_pictures(shape):
    if shape.shape_type == 6:
        for sub in list(shape.shapes):
            _drop_pictures(sub)
    elif shape.shape_type == 13:
        shape._element.getparent().remove(shape._element)


def shape_by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(f"{name} not found")


# ------------------------------------------------------------------- text API
def set_runs_in_place(shape, texts):
    """Replace the text of existing runs, preserving template formatting.
    Maps texts onto paragraphs that actually contain runs (skips empty ones)."""
    paras = [p for p in shape.text_frame.paragraphs if p.runs]
    assert len(paras) >= len(texts), f"{shape.name}: {len(paras)} run-paras < {len(texts)}"
    for para, text in zip(paras, texts):
        para.runs[0].text = text
        for r in para.runs[1:]:
            r.text = ""


def write_text(shape, blocks):
    """Rebuild a text frame. blocks = [(text, size_pt, bold, hexcolor), ...]."""
    tf = shape.text_frame
    tf.clear()
    first = True
    for text, size, bold, color in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        # strip any inherited numbering/bullet props from the template
        pPr = p._pPr if p._pPr is not None else p.get_or_add_pPr()
        for tag in ("buAutoNum", "buChar", "buNone"):
            for el in pPr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}"):
                pPr.remove(el)
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = text
        f = run.font
        f.name = "Arial"
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = RGBColor.from_string(color)


def title(slide, text):
    write_text(shape_by_name(slide, "TextBox 3"), [(text, 32, True, "FFFFFF")])


BLACK, DARK, WHITE, MINT = "000000", "1A1A1A", "FFFFFF", "17F1BD"
HEAD = "033529"

def bullets(*lines, size=18, head_first=False):
    out = []
    for i, ln in enumerate(lines):
        if ln.startswith("**"):
            out.append((ln.strip("*"), size + 2, True, HEAD))
        else:
            out.append(("• " + ln, size, False, DARK))
    return out


# ===================================================================== SLIDES
# --- 1 · Title (dup slide 1)
s, pics = dup_slide(0)
fix_pictures(s, pics)
set_runs_in_place(shape_by_name(s, "TextBox 6"), [
    "ICRISAT Data Hub",
    "Federation-First Data Sharing and an AI-Ready Catalogue for ICRISAT Teams",
    "\t\tJuly 2026",
])
set_runs_in_place(shape_by_name(s, "Rectangle 8"), [
    "Authors:",
    "Henok Desalegn, HaFAS Team",
])

# --- 2 · Challenge (dup slide 2)
s, _pics = dup_slide(1)
title(s, "The Challenge: Fragmented Team Data")
write_text(shape_by_name(s, "TextBox 4"), bullets(
    "**Current data-sharing situation**",
    "Valuable datasets sit scattered across teams, laptops, and shared drives — hard to find, easy to duplicate",
    "Sharing needs technical skills, so most data never leaves the team that produced it",
    "No common catalogue: nobody knows what exists, who owns it, or whether it is open, internal, or restricted",
    "**Consequences**",
    "Duplicated collection and processing effort across teams",
    "High-value compilations (e.g. DST nutrient-response trials) invisible beyond their owners",
    "Institutional memory lost when staff move on",
    "Institutional data unreachable by AI assistants and LLM-based tools",
    "**Need: zero-skill sharing, one catalogue, AI-ready access — without changing how teams work**",
    size=17))

# --- 3 · Solution (dup slide 3; right picture -> dashboard overview)
s, pics = dup_slide(2)
title(s, "Our Solution: The ICRISAT Data Hub")
write_text(shape_by_name(s, "TextBox 4"), bullets(
    "**Zero-skill upload — three ways in**",
    "Drop folder, drag-&-drop web page (token-protected link), or a tiny .meta.yaml sidecar",
    "**Automated pipeline**",
    "Reads CSV / Excel / Word, profiles every sheet, dedups by content hash, flags metadata gaps with a pre-filled review template",
    "**Federation-first**",
    "External data registered as YAML pointers — the hub points to where it lives (CHIRPS, AClimate)",
    "**AI-ready by design**",
    "MCP server lets any LLM search, inspect schemas, preview and query (read-only SQL), stdio or HTTP",
    "**Public dashboard**",
    "ICRISAT-branded Asset Explorer on GitHub Pages →",
    size=16))
fix_pictures(s, pics, {"Picture 8": str(SHOTS / "Screenshot 2026-07-29 141435.png")})

# --- 4 · Architecture (dup slide 4, two columns)
s, pics = dup_slide(3)
title(s, "System Architecture")
write_text(shape_by_name(s, "TextBox 4"), [("Technology Stack", 20, True, HEAD)])
write_text(shape_by_name(s, "TextBox 5"), bullets(
    "Upload: FastAPI web app (token-gated) + folder watcher (watchdog)",
    "Pipeline: Python — pandas, openpyxl, python-docx, DuckDB",
    "Catalogue: SQLite + exported assets.json (CDH-style)",
    "AI access: MCP server (MCP 2.0 SDK), stdio + streamable HTTP",
    "Dashboard: Chart.js + vanilla JS on GitHub Pages; auto-refresh via GitHub Actions",
    "Deployment: local, cloudflared/ngrok tunnel, or Docker on Railway/Render",
    size=15))
write_text(shape_by_name(s, "TextBox 7"), [("Data Flow", 20, True, HEAD)])
write_text(shape_by_name(s, "TextBox 8"), bullets(
    "Team drops a file (folder, web page, or with sidecar metadata)",
    "Pipeline reads and profiles every sheet; hash-dedups against the catalogue",
    "Asset registered in SQLite with status published or needs_review",
    "Catalogue exported to assets.json; dashboard data regenerated",
    "Dashboard updated on GitHub Pages (commit or GitHub Action)",
    "LLMs discover and query assets through the MCP server",
    size=15))
fix_pictures(s, pics)

# --- 5 · UI Overview (dup slide 5)
s, pics = dup_slide(4)
title(s, "User Interface: Public Dashboard — Portfolio Overview")
write_text(shape_by_name(s, "TextBox 4"), bullets(
    "**Coverage & insights**",
    "Team × file-type coverage matrix — click a cell to drill in",
    "Auto-generated insights: largest contributor, newest addition, review status",
    "Assets-per-team bars split by catalogue status",
    "**Charts**",
    "Domain mix, data access (open / internal / restricted), file types",
    "KPI strip: assets, teams, open-access share, published, needs review — clickable",
    "Faceted filters with per-value counts; every figure computed from the catalogue",
    size=15))
fix_pictures(s, pics, {"Picture 7": str(SHOTS / "Screenshot 2026-07-29 141435.png")})

# --- 6 · UI Explore (dup slide 7)
s, pics = dup_slide(6)
title(s, "User Interface: Explore the Catalogue")
write_text(shape_by_name(s, "TextBox 4"), bullets(
    "**Find any asset**",
    "Full-text search across titles, descriptions, column names, tags, and document text",
    "Filter by team, file type, status, access, and domain",
    "Sortable table; Export CSV of the current view",
    "**Inspect before you download**",
    "Row click opens a detail drawer: full metadata, provenance, and per-sheet schema",
    "Column names, data types, and sample values for every table",
    "Word documents show extracted text and outline",
    "Access badges make open vs internal vs restricted unmistakable",
    size=15))
fix_pictures(s, pics, {"Picture 8": str(SHOTS / "Screenshot 2026-07-29 141538.png")})

# --- 7 · UI Upload & Sources (dup slide 6; two screenshots in the picture slot)
s, _pics = dup_slide(5)
title(s, "User Interface: Upload Channels & Federated Sources")
write_text(shape_by_name(s, "TextBox 4"), bullets(
    "**Upload view**",
    "How data arrives: drop folder vs web form vs documented sidecar",
    "Metadata-completeness bar and most-missing-fields ranking",
    "Catalogue growth timeline",
    "**Sources view**",
    "Federated external datasets/APIs registered as YAML pointers",
    "Domain definitions from the controlled vocabulary",
    size=15))
# drop the template picture, stack the two screenshots in its slot
for shape in list(s.shapes):
    if shape.shape_type == 13:
        shape._element.getparent().remove(shape._element)
img = Image.open(SHOTS / "Screenshot 2026-07-29 141603.png")
aspect = img.width / img.height
EMU = 914400
slot_l, slot_t, slot_w, slot_h = 4.4, 1.3, 8.9, 6.2
each_h = (slot_h - 0.15) / 2
each_w = each_h * aspect
each_l = slot_l + (slot_w - each_w) / 2
s.shapes.add_picture(str(SHOTS / "Screenshot 2026-07-29 141603.png"),
                     int(each_l * EMU), int(slot_t * EMU), int(each_w * EMU), int(each_h * EMU))
s.shapes.add_picture(str(SHOTS / "Screenshot 2026-07-29 141631.png"),
                     int(each_l * EMU), int((slot_t + each_h + 0.15) * EMU), int(each_w * EMU), int(each_h * EMU))

# --- 8 · Core implementation (dup slide 8, code box)
s, _pics = dup_slide(7)
title(s, "Core Implementation (Ingest Pipeline + MCP)")
code = '''# src/hub/ingest/pipeline.py — one ingest path for every channel
def ingest_file(file_path, config, catalog, extra_metadata=None):
    digest = sha256_of(file_path)            # dedup by content hash
    if catalog.has_hash(digest):
        return {"ok": False, "reason": "duplicate"}
    tables  = read_file(file_path)           # CSV | Excel (multi-sheet) | Word
    meta    = infer_metadata(file_path, tables)      # filename, shape, columns
    meta.update(load_sidecar_metadata(file_path))    # optional .meta.yaml
    profile = profile_tables(tables, config)         # schema + sample values
    asset_id = catalog.add_asset(record, profile)    # SQLite + assets.json
    if incomplete: write_review_template(asset_id)   # pre-filled YAML

# MCP tools exposed to any LLM (stdio + streamable HTTP)
TOOLS = ["list_assets", "search_assets", "get_asset", "preview_asset",
         "query_asset",          # read-only DuckDB SQL over CSV/Excel
         "read_document",        # extracted Word text
         "list_sources",         # federated YAML pointers
         "get_catalog_summary"]  # portfolio statistics'''
tb = shape_by_name(s, "TextBox 5")
tf = tb.text_frame
tf.clear()
first = True
for line in code.split("\n"):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    run = p.add_run()
    run.text = line if line else " "
    f = run.font
    f.name = "Consolas"
    f.size = Pt(13)
    f.color.rgb = RGBColor.from_string("E4E7EC")

# --- 9 · Catalogue today (dup slide 2)
s, _pics = dup_slide(1)
title(s, "The Catalogue Today: First Assets Ingested")
write_text(shape_by_name(s, "TextBox 4"), bullets(
    "**One week in — the pipeline already runs on real team data**",
    "7 assets catalogued from 2 teams (5 CSV, 1 Excel, 1 Word); 5 published, 2 completing metadata review",
    "4 HaFAS DST nutrient-response compilations for Ethiopia (1990–2025):",
    "   Wheat — 18,973 georeferenced trial records with N/P/K rates, treatments, and grain yield",
    "   Tef — 17,696 records  ·  Maize — 16,028 records  ·  Sorghum — 6,613 records",
    "≈ 59,300 trial records now searchable, previewable, and queryable by people and LLMs alike",
    "**Governance made visible**",
    "Every asset carries team, owner, license, and access — open vs internal vs restricted is explicit, not assumed",
    "2 federated sources registered as YAML pointers (CHIRPS daily rainfall; AClimate Ethiopia advisory API)",
    "**Next: onboard remaining teams, deploy the permanent upload link (Railway), extend readers to GeoTIFF/NetCDF/parquet**",
    size=17))

# --- 10 · Impact (dup slide 9)
s, _pics = dup_slide(8)
title(s, "Expected Impact & Conclusion")
write_text(shape_by_name(s, "TextBox 4"), [("Expected Benefits", 20, True, HEAD)])
write_text(shape_by_name(s, "TextBox 5"), bullets(
    "Zero-friction sharing: any team member can publish a dataset in 30 seconds — no special skills",
    "One catalogue: what exists, who owns it, and its access terms are always known",
    "No more duplicate collection/processing; shared inputs registered once and reused",
    "Institutional data becomes AI-ready: LLM assistants answer questions directly from team data",
    "Clear open/internal/restricted governance supports data-policy compliance",
    size=16))
write_text(shape_by_name(s, "TextBox 7"), [("Key Innovations", 20, True, HEAD)])
write_text(shape_by_name(s, "TextBox 8"), bullets(
    "Federation-first: YAML pointers register external data without copying it",
    "Auto-infer + review-template metadata flow — gaps flagged, never blocking upload",
    "MCP server turns an institutional catalogue into LLM-consumable tools",
    "Deploy-anywhere: laptop, tunnel, or cloud container with the same codebase",
    "Aligned with the CGIAR Climate Data Hub asset-mapping approach and the Coalition of the Willing for soil and agronomy data sharing",
    size=16))
set_runs_in_place(shape_by_name(s, "TextBox 9"), [
    "The ICRISAT Data Hub turns scattered team files into a governed, AI-ready, federation-first catalogue — making ICRISAT's data as easy to share as it is to produce.",
])

# --- 11 · Citation & acknowledgements (dup slide 10)
s, pics = dup_slide(9)
fix_pictures(s, pics)
shape_by_name(s, "Rectangle 3").text_frame.clear()
tf = shape_by_name(s, "Rectangle 3").text_frame
citation_blocks = [
    ("Citation:", 16, True),
    ("Desalegn, H. 2026. ICRISAT Data Hub: Federation-First Data Sharing and an AI-Ready Catalogue for ICRISAT Teams. ICRISAT Working Paper.", 14, False),
    ("Acknowledgements", 16, True),
    ("The CGIAR Sustainable Farming Science Program forms a part of CGIAR's Research Portfolio, addressing key challenges in agri-food systems by fostering efficient production of nutritious foods and safeguarding the environment. This work was implemented by CGIAR researchers from ICRISAT in close partnership with the Ministry of Agriculture, the Ethiopian Institute of Agricultural Research, and Regional Agricultural Research Institutes. The hub design follows the CGIAR Climate Data Hub asset-mapping approach and supports the Coalition of the Willing for soil and agronomy data access, management, and sharing.", 12, False),
    ("We thank all funders who supported this research through their contributions to the CGIAR Trust Fund: https://www.cgiar.org/funders/", 12, False),
    ("Disclaimer", 16, True),
    ("This working paper has not been peer reviewed. Any opinions stated herein are those of the author(s) and do not necessarily reflect the policies or opinions of ICRISAT, donors, or partners.", 12, False),
]
first = True
for text, size, bold in citation_blocks:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(8)
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = "Arial"
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(DARK)

# --- 12 · License & keywords (dup slide 11)
s, pics = dup_slide(10)
fix_pictures(s, pics)
shape_by_name(s, "Rectangle 3").text_frame.clear()
tf = shape_by_name(s, "Rectangle 3").text_frame
license_blocks = [
    ("This publication is copyrighted by ICRISAT and licensed under a Creative Commons Attribution–NonCommercial 4.0 International License. Attribution must not suggest endorsement by ICRISAT or the author(s).", 12, False),
    ("Key Words:", 14, True),
    ("Data Hub; Federation; Data Catalogue; Model Context Protocol (MCP); AI-Ready Data; FAIR Data; Zero-Skill Upload", 12, False),
    ("Alignment:", 14, True),
    ("CGIAR Climate Data Hub asset-mapping approach · Coalition of the Willing for soil and agronomy data access, management and sharing · Ethiopian Ministry of Agriculture data-sharing policy", 12, False),
    ("©2026 ICRISAT", 12, False),
]
first = True
for text, size, bold in license_blocks:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(8)
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = "Arial"
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(DARK)

# --- 13 · Back cover (dup slide 12, graphics kept)
s, pics = dup_slide(11)
fix_pictures(s, pics)

# ------------------------------------------------- delete original template slides
sld_id_lst = prs.slides._sldIdLst
for _ in range(N_ORIG):
    rId = sld_id_lst[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    prs.part.drop_rel(rId)
    sld_id_lst.remove(sld_id_lst[0])

prs.save(str(OUT))
print(f"saved {OUT} with {len(prs.slides._sldIdLst)} slides")
